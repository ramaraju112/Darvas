#!/usr/bin/env python3
"""Darvas V6 research-only multi-box backtest. No broker orders or live portfolio accounting.
Pivot/box implementation is a transparent approximation, not a verified reproduction of Darvas's proprietary rules.
"""
import argparse, math, json, time
from pathlib import Path
import numpy as np
import pandas as pd
from darvas_fresh_v4 import normalize, hma, download

INITIAL=5000.0
ADD=1250.0


def as_date(d,i): return str(d.index[i].date())
def money(v): return round(float(v),2)
def event(events,ticker,pid,day,kind,price,shares,invested,stop,comment,**extra):
    events.append(dict(ticker=ticker,position_id=pid,date=day,event=kind,price=money(price) if price is not None else None,shares=shares,invested_total=money(invested),active_stop=money(stop) if stop is not None else None,commentary=comment,**extra))

def box_candidates(d):
    """Causal pivots, confirmed by 2 subsequent bars. At index i, only pivots through i-3 are used."""
    hi=d.High.to_numpy(); lo=d.Low.to_numpy(); n=len(d)
    tops=[]; bottoms=[]; output=[None]*n
    for i in range(n):
        j=i-3
        if j>=2:
            if hi[j]>max(hi[j-2:j]) and hi[j]>=max(hi[j+1:j+3]): tops.append(j)
            if lo[j]<min(lo[j-2:j]) and lo[j]<=min(lo[j+1:j+3]): bottoms.append(j)
        if i<60: continue
        # Prefer the latest still-valid top; do not reuse stale top after a new higher high.
        for top in reversed(tops):
            if top<i-120: break
            if np.max(hi[top+1:i])>hi[top]+1e-8: continue
            bottom=next((b for b in reversed(bottoms) if b>=top-50 and b<=i-3 and lo[b]<hi[top] and np.min(lo[b+1:i])>=lo[b]-1e-8),None)
            if bottom is not None and i-max(top,bottom)>=3:
                output[i]=(round(float(hi[top]),5),round(float(lo[bottom]),5),top,bottom)
                break
    return output

def run_one(ticker,frame,config=None):
    config=config or {}; d=normalize(frame)
    if len(d)<85: return [],[],[],[],[],dict(ticker=ticker,status='INSUFFICIENT_HISTORY',bars=len(d))
    n=len(d); H=d.High.to_numpy(); L=d.Low.to_numpy(); C=d.Close.to_numpy(); O=d.Open.to_numpy(); V=d.Volume.to_numpy()
    hm=hma(d.Close,38).to_numpy(); boxes=box_candidates(d)
    fee=float(config.get('fee_per_order',0)); slip=float(config.get('slippage_bps',0))/10000
    trades=[]; events=[]; boxlog=[]; adds=[]; stops=[]; current=None; pending=None; serial=0; last_setup=-1
    for i in range(60,n):
        day=as_date(d,i)
        # Pending orders are executed at today's open, before any intraday stop evaluation.
        if pending:
            p=pending; pending=None
            price=float(O[i])*(1+slip); qty=int(p['budget']//price) if price>0 else 0
            if qty==0 or (current is None and p['kind']=='ADD'):
                if p['kind']=='ADD':
                    adds.append(dict(ticker=ticker,position_id=p['pid'],box_number=p['box_no'],signal_date=p['signal_date'],execution_date=day,budget=ADD,shares=0,invested=0,status='UNFILLED',reason='NO_POSITION_OR_PRICE_EXCEEDS_BUDGET'))
                    event(events,ticker,p['pid'],day,'ADD UNFILLED',price,0,current['spent'] if current else 0,current['stop'] if current else None,'No whole share affordable or no open position')
                if p['kind']=='BUY':
                    event(events,ticker,f'{ticker}-UNFILLED',day,'BUY UNFILLED',price,0,0,p['low'],'No whole share affordable with $5,000 budget')
                continue
            if p['kind']=='BUY':
                serial+=1; pid=f'{ticker}-{serial:03d}'; stop=p['low']
                if price<=stop:
                    event(events,ticker,pid,day,'BUY SKIPPED',price,0,0,stop,'Opening fill price at/below planned stop; no entry')
                    continue
                current=dict(ticker=ticker,position_id=pid,entry_date=day,entry_index=i,entry_price=price,shares=qty,spent=qty*price+fee,fees=fee,stop=stop,box_high=p['high'],box_low=p['low'],box_no=1,addon_count=0,last_add_box=1,ht1_date='',ht2_date='',ht1=p['high']+(p['high']-p['low']),ht2=p['high']+2*(p['high']-p['low']),stop_effective_date=day)
                event(events,ticker,pid,day,'BUY',price,qty,current['spent'],stop,f'Initial $5,000 budget: bought {qty} whole shares at next-session open; box 1 low is initial stop')
                stops.append(dict(ticker=ticker,position_id=pid,date=day,previous_stop=None,new_stop=stop,reason='INITIAL_BOX_LOW'))
                boxlog.append(dict(ticker=ticker,position_id=pid,box_number=1,signal_date=p['signal_date'],box_high=p['high'],box_low=p['low'],breakout_date=p['signal_date'],action='INITIAL_BUY'))
            elif p['kind']=='ADD' and current and current['position_id']==p['pid']:
                current['shares']+=qty; current['spent']+=qty*price+fee; current['fees']+=fee; current['addon_count']+=1
                adds.append(dict(ticker=ticker,position_id=p['pid'],box_number=p['box_no'],signal_date=p['signal_date'],execution_date=day,budget=ADD,shares=qty,invested=money(qty*price+fee),status='FILLED',reason='HIGHER_BOX_BREAKOUT'))
                event(events,ticker,p['pid'],day,'ADD',price,qty,current['spent'],current['stop'],f'Automatic $1,250 higher-box add-on; spent ${qty*price+fee:,.2f}; cumulative invested ${current["spent"]:,.2f}')
        if current:
            p=current; stop=p['stop']; reason=None; exit_price=None
            if i>p['entry_index'] and O[i]<=stop:
                reason='STOP_GAP'; exit_price=float(O[i])*(1-slip)
            elif L[i]<=stop:
                reason='STOP_INTRADAY'; exit_price=float(stop)*(1-slip)
            if reason:
                proceeds=p['shares']*exit_price-fee; pnl=proceeds-p['spent']
                event(events,ticker,p['position_id'],day,'EXIT',exit_price,-p['shares'],p['spent'],stop,f'{reason}: sold {p["shares"]} shares; invested ${p["spent"]:,.2f}, proceeds ${proceeds:,.2f}, realized P/L ${pnl:+,.2f}')
                trades.append(dict(**p,status='CLOSED',exit_date=day,exit_price=money(exit_price),exit_reason=reason,proceeds=money(proceeds),realized_pl=money(pnl),unrealized_pl=None,return_pct=round(100*pnl/p['spent'],2),calendar_days=(d.index[i]-d.index[p['entry_index']]).days,trading_sessions=i-p['entry_index']+1,final_mark=money(exit_price)))
                current=None; pending=None; last_setup=i
                continue
            # HT milestones never close the trade; stop-first daily convention.
            if not p['ht1_date'] and H[i]>=p['ht1']:
                p['ht1_date']=day;event(events,ticker,p['position_id'],day,'HT1 MILESTONE',p['ht1'],0,p['spent'],p['stop'],'HT1 observed; continue holding')
            if not p['ht2_date'] and H[i]>=p['ht2']:
                p['ht2_date']=day;event(events,ticker,p['position_id'],day,'HT2 MILESTONE',p['ht2'],0,p['spent'],p['stop'],'HT2 observed; NOT a mandatory exit')
        b=boxes[i]
        if b is None: continue
        bh,bl,top,bottom=b
        # Close breakout must cross box high today, from at/below it yesterday.
        breakout=(C[i]>bh and C[i-1]<=bh)
        if current:
            p=current
            higher=bh>p['box_high']+1e-8 and bl>p['box_low']+1e-8
            if higher:
                old=p['stop'];p['box_high']=bh;p['box_low']=bl;p['box_no']+=1
                boxlog.append(dict(ticker=ticker,position_id=p['position_id'],box_number=p['box_no'],signal_date=day,box_high=bh,box_low=bl,breakout_date=day if breakout else '',action='HIGHER_BOX_CONFIRMED'))
                event(events,ticker,p['position_id'],day,'HIGHER BOX',C[i],0,p['spent'],old,f'Confirmed higher box {p["box_no"]}: high ${bh:.2f}, low ${bl:.2f}')
                # A stop ratchet confirmed at close is effective NEXT session, never earlier.
                if bl>old and bl<C[i]:
                    p['stop']=bl;stops.append(dict(ticker=ticker,position_id=p['position_id'],date=day,previous_stop=old,new_stop=bl,reason='HIGHER_BOX_CONFIRMED_EFFECTIVE_NEXT_SESSION'))
                    event(events,ticker,p['position_id'],day,'STOP RAISED',C[i],0,p['spent'],bl,f'Protective stop raised ${old:.2f} → ${bl:.2f}; effective from next trading session')
            # The new box can be confirmed before its breakout; retain it until a later crossover.
            if breakout and p['box_no']>p['last_add_box']:
                p['last_add_box']=p['box_no']
                if i<n-1:
                    pending=dict(kind='ADD',pid=p['position_id'],box_no=p['box_no'],signal_date=day,budget=ADD)
                    event(events,ticker,p['position_id'],day,'ADD SIGNAL',C[i],0,p['spent'],p['stop'],'Confirmed higher-box close breakout; automatic $1,250 budget order for next open (no volume/HMA add-on veto)')
                else:
                    adds.append(dict(ticker=ticker,position_id=p['position_id'],box_number=p['box_no'],signal_date=day,execution_date='',budget=ADD,shares=0,invested=0,status='UNFILLED',reason='NO_NEXT_SESSION'))
            continue
        # Fresh initial setup, using HMA/volume/extension; do not enter after last session.
        if i<=last_setup or i==n-1 or not breakout: continue
        volbase=np.mean(V[i-20:i]);ratio=V[i]/volbase if volbase>0 else 0
        rising=np.isfinite(hm[i]) and np.isfinite(hm[i-1]) and hm[i]>=hm[i-1]
        ext=C[i]/bh-1
        if ratio>=float(config.get('initial_volume_ratio',1.5)) and rising and ext<=float(config.get('initial_max_extension',.03)):
            pending=dict(kind='BUY',budget=INITIAL,signal_date=day,high=bh,low=bl)
            last_setup=i
    if current:
        p=current; last=n-1; mark=float(C[last]); unreal=p['shares']*mark-p['spent']
        event(events,ticker,p['position_id'],as_date(d,last),'OPEN AT TEST END',mark,0,p['spent'],p['stop'],f'No simulated sale; market value ${p["shares"]*mark:,.2f}; unrealized P/L ${unreal:+,.2f}')
        trades.append(dict(**p,status='OPEN_AT_TEST_END',exit_date='',exit_price=None,exit_reason='CENSORED',proceeds=None,realized_pl=None,unrealized_pl=money(unreal),return_pct=round(100*unreal/p['spent'],2),calendar_days=(d.index[last]-d.index[p['entry_index']]).days,trading_sessions=last-p['entry_index']+1,final_mark=money(mark)))
    return trades,events,boxlog,adds,stops,dict(ticker=ticker,status='OK',bars=n,start=as_date(d,0),end=as_date(d,n-1))

def excel(path, data, settings):
    """Write the ten-tab research report with public openpyxl."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from datetime import date, datetime

    wb = Workbook()
    wb.remove(wb.active)
    names = ['SUMMARY', 'TRADE LEDGER', 'TRADE JOURNEY', 'BOX HISTORY',
             'ADD-ON HISTORY', 'STOP HISTORY', 'MILESTONES', 'OPEN AT END',
             'RULES & AUDIT', 'DATA COVERAGE']
    trade, journey, box, addon, stop, coverage = data
    closed = [t for t in trade if t['status'] == 'CLOSED']
    opened = [t for t in trade if t['status'] != 'CLOSED']
    summary = [
        dict(metric='Tickers checked', value=len(coverage)),
        dict(metric='Completed positions', value=len(closed)),
        dict(metric='Open at test end', value=len(opened)),
        dict(metric='Realized net P/L USD', value=money(sum(t['realized_pl'] for t in closed))),
        dict(metric='Unrealized P/L USD', value=money(sum(t['unrealized_pl'] for t in opened))),
        dict(metric='Add-on fills', value=sum(a['status'] == 'FILLED' for a in addon)),
        dict(metric='Stop exits', value=sum(t['exit_reason'].startswith('STOP') for t in closed)),
        dict(metric='Portfolio equity curve', value='NOT CALCULATED: isolated position replay')]
    rules = [
        dict(rule='Initial budget per position', value=INITIAL),
        dict(rule='Each higher-box add budget', value=ADD),
        dict(rule='Initial entry', value='Next open after qualified daily breakout'),
        dict(rule='Add-on entry', value='Next open after higher-box close breakout; no extra filters'),
        dict(rule='Stop', value='Higher confirmed box low, ratchet up only, effective next bar'),
        dict(rule='Stop gaps', value='Open below stop exits at open minus slippage'),
        dict(rule='HT1 / HT2', value='Milestones, never terminal exits'),
        dict(rule='Fractional shares', value='NO'),
        dict(rule='Portfolio cash limits', value='NOT MODELED; independent per-position capital'),
        dict(rule='Same-bar ambiguity', value='Stop-first; OHLC cannot sequence events'),
        dict(rule='Survivorship bias', value='Present-day symbols are not point-in-time universe'),
        dict(rule='Adjustment', value='Unadjusted OHLC; splits/dividends can distort results'),
        dict(rule='Data source', value='Yahoo via yfinance; best effort'),
        dict(rule='Commission per order', value=settings.get('fee_per_order', 0)),
        dict(rule='Slippage bps', value=settings.get('slippage_bps', 0))]
    mapping = [summary, trade, journey, box, addon, stop,
               [dict(ticker=t['ticker'], position_id=t['position_id'],
                     ht1_first_touch=t['ht1_date'], ht2_first_touch=t['ht2_date']) for t in trade],
               opened, rules, coverage]
    for name, rows in zip(names, mapping):
        ws = wb.create_sheet(name)
        headers = list(dict.fromkeys(key for row in rows for key in row)) if rows else ['message']
        ws.append(headers)
        if rows:
            for row in rows:
                ws.append([v if isinstance(v, (str, int, float, bool, date, datetime)) or v is None
                           else str(v) for v in (row.get(k) for k in headers)])
        else:
            ws.append(['No records'])
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = PatternFill('solid', fgColor='17365D')
            cell.font = Font(color='FFFFFF', bold=True)
            cell.alignment = Alignment(wrap_text=True)
        for idx, header in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(idx)].width = min(48, max(14, len(str(header)) + 3))
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))

def main():
    a=argparse.ArgumentParser();a.add_argument('--symbols',default='symbols.txt');a.add_argument('--period',default='5y');a.add_argument('--output',default='darvas_v6_backtest_report.xlsx');a.add_argument('--limit',type=int,default=0);a.add_argument('--fee',type=float,default=0);a.add_argument('--slippage-bps',type=float,default=0);a.add_argument('--data-dir',default='');args=a.parse_args()
    symbols=list(dict.fromkeys(x.strip().upper() for x in Path(args.symbols).read_text().splitlines() if x.strip() and not x.startswith('#')))
    if args.limit:symbols=symbols[:args.limit]
    allrows=[[],[],[],[],[],[]];cfg=dict(fee_per_order=args.fee,slippage_bps=args.slippage_bps)
    for k,t in enumerate(symbols,1):
        try:
            if args.data_dir:
                path=Path(args.data_dir)/f'{t}.csv'; d=pd.read_csv(path,index_col=0,parse_dates=True)
            else:d=download(t,args.period)
            result=run_one(t,d,cfg)
            for target,part in zip(allrows,result):target.extend(part if isinstance(part,list) else [part])
            print(f'[{k}/{len(symbols)}] {t}: {len(result[0])} positions, {len(result[3])} add-on events',flush=True)
        except Exception as exc:
            allrows[5].append(dict(ticker=t,status='ERROR',detail=f'{type(exc).__name__}: {str(exc)[:180]}'))
            print(f'[{k}/{len(symbols)}] {t}: ERROR {exc}',flush=True)
    excel(args.output,allrows,cfg)
    print('Saved',args.output)
if __name__=='__main__':main()
