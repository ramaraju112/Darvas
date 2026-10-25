#!/usr/bin/env python3
"""Rank V5 diagnostic qualified signals for MANUAL research; no trades or live quotes.

Usage: python darvas_buy_ranking.py --input darvas_v5_scan_report.xlsx
Creates BUY RANKING and RANKING GUIDE sheets in the existing workbook.
"""
import argparse
from datetime import date, datetime
from math import isfinite
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = ['Rank','Ticker','Signal Date','Close','Box High','Stop','Volume x',
           'Extension %','Stop Risk %','HT2 Reward/Risk','Breakout (25)',
           'Volume (20)','Trend (20)','Entry (20)','Risk (15)',
           'Total / 100','Classification','Review Notes']

def num(v):
    try:
        x=float(v)
        return x if isfinite(x) else None
    except (ValueError,TypeError):
        return None

def yes(v):
    return v is True or str(v).strip().lower() in ('true','yes','1')

def evaluate(d, latest):
    close, high, stop, vol = (num(d.get(k)) for k in ('close','box_high','stop','volume_ratio'))
    hma, prev, ht2 = (num(d.get(k)) for k in ('hma38','hma38_prev','ht2'))
    dt=d.get('date')
    if isinstance(dt,datetime): dt=dt.date()
    elif not isinstance(dt,date):
        try: dt=date.fromisoformat(str(dt)[:10])
        except ValueError: dt=None
    reasons=[]
    if str(d.get('status','')).strip().upper()!='QUALIFIED': reasons.append('Not qualified')
    stale = dt is None or (latest is not None and dt != latest)
    if stale: reasons.append('Historical signal; recheck today')
    if close is None or close<=0 or high is None or high<=0: reasons.append('Missing price/box')
    if stop is None or close is None or not (0<stop<close): reasons.append('Invalid stop')
    if vol is None or vol<1.5: reasons.append('Volume below 1.5x')
    if not yes(d.get('fresh_cross')): reasons.append('Not fresh breakout')
    if not yes(d.get('hma_pass')) or hma is None or prev is None or hma<=prev: reasons.append('HMA38 not rising')
    ext=(close/high-1)*100 if close and high and high>0 else None
    risk=(close-stop)/close*100 if close and stop and close>0 else None
    rr=(ht2-close)/(close-stop) if ht2 and close and stop and close>stop else None
    if ext is not None and ext>3: reasons.append('Extended >3% above box')
    if risk is not None and risk>8: reasons.append('Stop risk >8%')
    if ext is not None and ext<0: reasons.append('Below box high')
    # Explicit proposed, unvalidated scoring formula; gates override score.
    breakout=25 if yes(d.get('fresh_cross')) and close and high and close>high else 0
    volume=round(20*min(1,max(0,(vol or 0)/3)),1)
    trend=20 if hma is not None and prev is not None and hma>prev and yes(d.get('hma_pass')) else 0
    entry=round(20*max(0,1-max(0,ext or 0)/5),1) if ext is not None and ext>=0 else 0
    risk_score=round(15*max(0,1-max(0,risk or 0)/12),1) if risk is not None and risk>0 else 0
    score=round(breakout+volume+trend+entry+risk_score,1)
    # Separate signal age from failed trading rules. A stale signal can have a high score.
    # Reject has priority if any rule other than recency fails.
    hard_failures = [r for r in reasons if r != 'Historical signal; recheck today']
    if hard_failures:
        classification = 'REJECT - FAILED RULES'
    elif stale:
        classification = 'STALE SIGNAL'
    elif score >= 70:
        classification = 'RESEARCH CANDIDATE'
    else:
        classification = 'WATCH / RECHECK'
    if rr is not None and rr<1.5: reasons.append('HT2 reward/risk <1.5 (informational)')
    return [None,d.get('ticker'),dt,close,high,stop,vol,ext/100 if ext is not None else None,
            risk/100 if risk is not None else None,rr,breakout,volume,trend,entry,risk_score,
            score,classification,'; '.join(reasons) or 'Confirm next-session price, earnings, liquidity and stop']

def run(path):
    wb=load_workbook(path)
    if 'BUY - Review' not in wb: raise ValueError('Missing BUY - Review worksheet')
    source=wb['BUY - Review']; records=list(source.values)
    if not records: raise ValueError('Empty BUY - Review worksheet')
    headers=[str(x).strip() for x in records[0]]
    required={'ticker','date','status','close','box_high','stop','volume_ratio','fresh_cross','hma_pass','hma38','hma38_prev'}
    missing=required-set(headers)
    if missing: raise ValueError('Missing columns: '+', '.join(sorted(missing)))
    # Latest date across all scan observations, not only qualified signals.
    scan=wb['SCAN - All 400'] if 'SCAN - All 400' in wb else source
    scanhead=[str(x).strip() for x in next(scan.values)]
    dateidx=scanhead.index('date')
    dates=[]
    for r in list(scan.values)[1:]:
        v=r[dateidx]
        try: dates.append(v.date() if isinstance(v,datetime) else v if isinstance(v,date) else date.fromisoformat(str(v)[:10]))
        except (ValueError,TypeError): pass
    latest=max(dates) if dates else None
    scored=[evaluate(dict(zip(headers,row)),latest) for row in records[1:] if any(v is not None for v in row)]
    priority={'RESEARCH CANDIDATE':0,'WATCH / RECHECK':1,'STALE SIGNAL':2,'REJECT - FAILED RULES':3}
    scored.sort(key=lambda r:(priority.get(r[16],4),-(r[15] or 0),str(r[1])))
    for name in ('BUY RANKING','RANKING GUIDE'):
        if name in wb: del wb[name]
    ws=wb.create_sheet('BUY RANKING',0)
    ws.append(COLUMNS)
    for i,row in enumerate(scored,1): row[0]=i;ws.append(row)
    ws.freeze_panes='D2';ws.auto_filter.ref=ws.dimensions
    for c in ws[1]:
        c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17365D');c.alignment=Alignment(wrap_text=True)
    ws.row_dimensions[1].height=35
    for col in range(1,len(COLUMNS)+1):
        ws.column_dimensions[get_column_letter(col)].width= 50 if col==18 else 25 if col==17 else 18
    for row in ws.iter_rows(min_row=2):
        row[2].number_format='yyyy-mm-dd'
        for i in (3,4,5): row[i].number_format='$#,##0.00'
        for i in (7,8): row[i].number_format='0.0%'
        row[9].number_format='0.00'
        row[6].number_format='0.00"x"'
        if row[16].value=='RESEARCH CANDIDATE': row[16].fill=PatternFill('solid',fgColor='DDF2E2')
        elif row[16].value=='STALE SIGNAL': row[16].fill=PatternFill('solid',fgColor='FFF2CC')
        elif row[16].value=='REJECT - FAILED RULES': row[16].fill=PatternFill('solid',fgColor='FCE8E6')
    guide=wb.create_sheet('RANKING GUIDE')
    notes=[('Purpose','Prioritize V5 QUALIFIED diagnostic observations for MANUAL RESEARCH only; NOT executable buy orders.'),
      ('Latest scan date',str(latest)),('RESEARCH CANDIDATE','Latest scan date, all hard rules pass, score >=70; research only, not a buy order.'),
      ('WATCH / RECHECK','Latest scan date and all hard rules pass, but score below 70; reassess.'),
      ('STALE SIGNAL','Signal date is older than latest date (or missing); other hard rules pass. Recheck for a new breakout.'),
      ('REJECT - FAILED RULES','One or more non-date hard rules fail. Reasons shown in Review Notes; may also be stale.'),
      ('Score weights','Fresh breakout 25; volume 20; HMA38 rising 20; extension 20; stop distance 15.'),
      ('Breakout score','25 if fresh_cross and close > box high; otherwise 0.'),
      ('Volume score','20 * min(volume_ratio / 3, 1).'),('Trend score','20 if HMA38 > previous HMA38 and hma_pass; otherwise 0.'),
      ('Entry score','20 * max(0, 1 - extension_percent / 5), for nonnegative extension.'),
      ('Risk score','15 * max(0, 1 - stop_risk_percent / 12), for valid positive stop risk.'),
      ('Hard filters','Qualified, latest date, fresh cross, rising HMA, volume >=1.5x, extension <=3%, stop risk <=8%, valid stop.'),
      ('HT2 reward/risk','HT2 upside divided by entry-to-stop downside; below 1.5 is flagged for review, not a hard gate. HT2 is a box projection, not a forecast.'),
      ('Price convention','Uses historical signal-day CLOSE as hypothetical entry; NOT a current quote or guaranteed fill.'),
      ('Risk','Gap risk, stop slippage, commissions, corporate actions and earnings can change outcomes.'),
      ('Validation','Scoring thresholds are experimental; validate using out-of-sample historical backtests before reliance.'),
      ('Original data','Original SCAN - All 400, BUY - Review and READ ME sheets are preserved.')]
    for x in notes: guide.append(x)
    guide.column_dimensions['A'].width=25;guide.column_dimensions['B'].width=115
    for c in guide['A']: c.font=Font(bold=True)
    wb.save(path)
    print(f'Updated {path}: {len(scored)} ranked observations; latest scan {latest}; '+str(sum(r[16]=='RESEARCH CANDIDATE' for r in scored))+' research candidates')

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--input',default='darvas_v5_scan_report.xlsx');a=p.parse_args();run(Path(a.input))
