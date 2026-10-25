#!/usr/bin/env python3
"""Track V5-qualified Darvas signals through subsequent completed daily bars.

This is SIGNAL performance, not an executable V6 backtest. The registry is
persisted between GitHub Actions runs and never assumes a position was filled.
"""
import argparse
import csv
from datetime import date, datetime, timedelta
from pathlib import Path
import time

import pandas as pd
import yfinance as yf
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FIELDS = ['ticker', 'signal_date', 'breakout_close', 'original_stop', 'box_high']
HEADERS = ['Ticker','Signal Date','Breakout Close','Original Stop','Box High',
           'Latest Close Date','Latest Close','Change $','Change %',
           'Trading Days After','Highest High After','Maximum Gain %',
           'Lowest Low After','Maximum Decline %','Original Stop Touched?',
           'First Stop Touch Date','Signal Status','Data Notes']

def day(v):
    if isinstance(v, datetime): return v.date()
    if isinstance(v, date): return v
    return date.fromisoformat(str(v)[:10])

def number(v):
    try: return float(v)
    except (TypeError,ValueError): return None

def read_registry(path):
    result={}
    if path.exists():
        with path.open(newline='',encoding='utf-8') as f:
            for row in csv.DictReader(f):
                try:
                    k=(row['ticker'].strip().upper(),day(row['signal_date']).isoformat())
                    result[k]={x:row.get(x,'') for x in FIELDS}
                except (KeyError,ValueError): continue
    return result

def ingest_signals(wb, registry):
    if 'BUY - Review' not in wb:
        raise ValueError('Missing BUY - Review worksheet')
    rows=iter(wb['BUY - Review'].values)
    headers=[str(h).strip() for h in next(rows)]
    required={'ticker','date','status','close','stop','box_high'}
    if required-set(headers): raise ValueError('Missing BUY - Review columns: '+str(required-set(headers)))
    for values in rows:
        r=dict(zip(headers,values))
        if str(r.get('status','')).upper()!='QUALIFIED': continue
        ticker=str(r['ticker']).strip().upper()
        close=number(r['close']); stop=number(r['stop']); box=number(r['box_high'])
        if not ticker or not close or close<=0: continue
        try: d=day(r['date'])
        except (ValueError,TypeError): continue
        k=(ticker,d.isoformat())
        if k not in registry:
            registry[k]={'ticker':ticker,'signal_date':d.isoformat(),
                         'breakout_close':str(close),
                         'original_stop':str(stop) if stop is not None else '',
                         'box_high':str(box) if box is not None else ''}

def save_registry(path, registry):
    path.parent.mkdir(parents=True,exist_ok=True)
    tmp=path.with_suffix('.tmp')
    with tmp.open('w',newline='',encoding='utf-8') as f:
        w=csv.DictWriter(f,fieldnames=FIELDS)
        w.writeheader()
        for key in sorted(registry,key=lambda x:(x[1],x[0])): w.writerow(registry[key])
    tmp.replace(path)

def fetch_history(ticker, start):
    # Yahoo daily bars; auto_adjust=False keeps raw price convention consistent
    # with most unadjusted signal-day scanner outputs. Verify source consistency.
    last_error=''
    for attempt in range(3):
        try:
            hist=yf.Ticker(ticker).history(start=(start-timedelta(days=5)).isoformat(),
                     end=(date.today()+timedelta(days=1)).isoformat(),
                     interval='1d',auto_adjust=False,actions=False)
            if hist is None or hist.empty: return None,'No Yahoo daily bars'
            hist=hist[['High','Low','Close']].copy()
            hist.index=pd.to_datetime(hist.index).date
            hist=hist[~pd.Index(hist.index).duplicated(keep='last')].sort_index()
            hist=hist.dropna(subset=['High','Low','Close'])
            return hist,None if not hist.empty else 'No complete daily bars'
        except Exception as exc:
            last_error=str(exc)[:160]
            time.sleep(1+attempt)
    return None,'Yahoo download failed: '+last_error

def evaluate(r,hist,error):
    ticker=r['ticker']; sd=day(r['signal_date']); entry=number(r['breakout_close'])
    stop=number(r['original_stop']); box=number(r['box_high'])
    prefix=[ticker,sd,entry,stop,box]
    if hist is None or hist.empty:
        return prefix+[None]*11+['DATA MISSING',error or 'No bars']
    # Exclude signal day from stop and extrema: no same-day lookahead.
    after=hist.loc[[d for d in hist.index if d>sd]]
    if after.empty:
        return prefix+[None]*11+['NO LATER BAR','Signal is too recent or no later completed daily bar']
    last_date=after.index[-1]; last=float(after.iloc[-1]['Close'])
    hi=float(after['High'].max()); lo=float(after['Low'].min())
    touch=bool(stop is not None and (after['Low']<=stop).any())
    first_touch=next((d for d,v in after['Low'].items() if stop is not None and v<=stop),None)
    change=last-entry; pct=change/entry
    status='STOP TOUCHED' if touch else 'ABOVE SIGNAL' if last>entry else 'BELOW SIGNAL' if last<entry else 'UNCHANGED'
    note='Signal-close comparison only; stop touch uses daily LOW, not actual fill'
    if sd not in hist.index: note+='; original signal bar absent from downloaded data'
    return prefix+[last_date,last,change,pct,len(after),hi,(hi-entry)/entry,lo,(lo-entry)/entry,
                   'YES' if touch else 'NO',first_touch,status,note]

def build_sheet(wb,rows):
    for name in ['BREAKOUT TRACKER','PERFORMANCE SUMMARY']:
        if name in wb: del wb[name]
    ws=wb.create_sheet('BREAKOUT TRACKER',1)
    ws.append(HEADERS)
    for r in rows: ws.append(r)
    ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
    for i in range(1,len(HEADERS)+1):
        ws.column_dimensions[get_column_letter(i)].width=38 if i==18 else 23 if i in (15,16,17) else 19
    ws.row_dimensions[1].height=32
    for c in ws[1]:
        c.fill=PatternFill('solid',fgColor='17365D');c.font=Font(color='FFFFFF',bold=True)
        c.alignment=Alignment(wrap_text=True)
    for r in ws.iter_rows(min_row=2):
        for i in (1,5,15):r[i].number_format='yyyy-mm-dd'
        for i in (2,3,4,6,7,10,12):r[i].number_format='$#,##0.00'
        for i in (8,11,13):r[i].number_format='0.0%'
        status=r[16].value
        if status=='ABOVE SIGNAL':r[16].fill=PatternFill('solid',fgColor='DDF2E2')
        if status=='STOP TOUCHED':r[16].fill=PatternFill('solid',fgColor='FCE8E6')
    sm=wb.create_sheet('PERFORMANCE SUMMARY',2)
    counts={x:sum(r[16]==x for r in rows) for x in ['ABOVE SIGNAL','BELOW SIGNAL','UNCHANGED','STOP TOUCHED','NO LATER BAR','DATA MISSING']}
    valid=[r for r in rows if isinstance(r[8],float)]
    avg=sum(r[8] for r in valid)/len(valid) if valid else None
    sm.append(['Metric','Value'])
    for x,y in [('Tracked historical signals',len(rows)),('Signals with later close',len(valid)),
                ('Average change since signal',avg),*counts.items(),
                ('Interpretation','Signal price performance only; not realized P/L or a trade backtest.'),
                ('Stop convention','STOP TOUCHED if any subsequent daily low <= original suggested stop.'),
                ('Gap/slippage','Stop-touch does not guarantee an exit at the stop price.'),
                ('Data freshness','Uses Yahoo Finance completed daily bars; missing data flagged.'),
                ('Data coverage','Historical signals accumulate only from scans after registry first enabled, plus those in current input workbook.')]:sm.append([x,y])
    sm['B4'].number_format='0.0%'
    sm.column_dimensions['A'].width=34;sm.column_dimensions['B'].width=115
    for c in sm[1]:c.fill=PatternFill('solid',fgColor='17365D');c.font=Font(color='FFFFFF',bold=True)
    for c in sm['A']:c.font=Font(bold=True)

def run(args):
    path=Path(args.input);registry_path=Path(args.registry)
    wb=load_workbook(path)
    registry=read_registry(registry_path)
    ingest_signals(wb,registry)
    if not registry: raise ValueError('No qualified signals found to track')
    save_registry(registry_path,registry)
    earliest={}
    for t,d in registry: earliest[t]=min(earliest.get(t,d),d)
    histories={}
    for i,(ticker,start) in enumerate(sorted(earliest.items()),1):
        histories[ticker]=fetch_history(ticker,date.fromisoformat(start))
        print(f'History {i}/{len(earliest)}: {ticker} — {histories[ticker][1] or "OK"}',flush=True)
    rows=[]
    for key in sorted(registry,key=lambda k:(k[1],k[0]),reverse=True):
        h,err=histories[key[0]]
        rows.append(evaluate(registry[key],h,err))
    build_sheet(wb,rows)
    wb.save(path)
    print(f'Wrote {len(rows)} signal histories into {path}; registry: {registry_path}')

if __name__=='__main__':
    p=argparse.ArgumentParser()
    p.add_argument('--input',default='darvas_v5_scan_report.xlsx')
    p.add_argument('--registry',default='v6_breakout_signals.csv')
    run(p.parse_args())
